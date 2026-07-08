#!/usr/bin/env python3
"""Generate MainFrame_Devin_Demo.pptx for the MainFramePOC modernization demo.

Reproducible: run `python3 tools/demo/gen_pptx.py` from the repo root.
Requires the diagram PNGs under assets/ppt-images/ (run gen_images.py first).

Branding: only a user-provided pasted logo image (Wells Fargo + Mphasis
"The Next Applied") is used, on the title slide and a footer on every slide.
No brand assets are fabricated or downloaded. If the logo file is absent, a
clearly-marked placeholder is inserted and a message tells you exactly where to
drop the logo file. Wells Fargo hex values (red #D71E28, gold #FFCD41) are
approximations — replace with the official brand guide if available.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
IMG = os.path.join(ROOT, "assets", "ppt-images")
LOGO_DIR = os.path.join(ROOT, "assets", "logo")
# Preferred logo file names (any one of these will be picked up automatically)
LOGO_CANDIDATES = [
    "wells-fargo-mphasis.png", "logo.png", "wf-mphasis-logo.png",
    "wells-fargo-mphasis.jpg", "logo.jpg",
]
OUT_PPTX = os.path.join(ROOT, "MainFrame_Devin_Demo.pptx")

# ---- Palette (Wells Fargo + Mphasis corporate; hex are approximations) ----
WF_RED      = RGBColor(0xD7, 0x1E, 0x28)
WF_DARK_RED = RGBColor(0xA5, 0x12, 0x1B)
WF_GOLD     = RGBColor(0xFF, 0xCD, 0x41)
INK         = RGBColor(0x1F, 0x29, 0x33)
SLATE       = RGBColor(0x3E, 0x4C, 0x59)
STEEL       = RGBColor(0x52, 0x60, 0x6D)
LIGHT       = RGBColor(0xF5, 0xF7, 0xFA)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BORDER      = RGBColor(0xCB, 0xD2, 0xD9)
GREEN       = RGBColor(0x2F, 0x85, 0x5A)
BLUE        = RGBColor(0x2B, 0x6C, 0xB0)
GREY_BADGE  = RGBColor(0x8A, 0x94, 0x9E)

SW, SH = Inches(13.333), Inches(7.5)


def find_logo():
    for name in LOGO_CANDIDATES:
        p = os.path.join(LOGO_DIR, name)
        if os.path.exists(p):
            return p
    return None


LOGO = find_logo()


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri",
             italic=False, line_spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = font
    return tb


def add_bullets(slide, x, y, w, h, items, size=16, color=INK, bullet_color=WF_RED,
                gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r0 = p.add_run(); r0.text = "\u25A0  "
        r0.font.size = Pt(size); r0.font.color.rgb = bullet_color
        r0.font.bold = True; r0.font.name = "Calibri"
        r = p.add_run(); r.text = it
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
    return tb


def rrect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def shape_text(sp, text, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = "Calibri"


def place_logo(slide, x, y, w, h, placeholder_label="LOGO", compact=False):
    """Insert the pasted logo, or a clearly-marked placeholder box."""
    if LOGO:
        try:
            iw, ih = Image.open(LOGO).size
            ar = iw / ih
            # fit inside (w, h) preserving aspect ratio
            if w / h > ar:
                nh = h; nw = int(h * ar)
            else:
                nw = w; nh = int(w / ar)
            slide.shapes.add_picture(LOGO, int(x + (w - nw) / 2),
                                     int(y + (h - nh) / 2), width=nw, height=nh)
            return
        except Exception:
            pass
    # placeholder
    sp = rrect(slide, x, y, w, h, fill=LIGHT, line=WF_DARK_RED, line_w=1.25)
    if compact:
        shape_text(sp, "[ LOGO ]", size=9, color=WF_DARK_RED, bold=True)
    else:
        shape_text(sp, f"[ {placeholder_label} ]\nWells Fargo + Mphasis\n"
                   "The Next Applied\n(drop logo here)", size=9,
                   color=WF_DARK_RED, bold=True)


def add_footer(slide, page_no=None):
    # thin footer strip
    bar = rrect(slide, 0, Inches(7.12), SW, Inches(0.38), fill=LIGHT, line=None)
    add_text(slide, Inches(0.35), Inches(7.14), Inches(7.5), Inches(0.32),
             "Accelerating Mainframe Modernization with Devin", size=9,
             color=STEEL, anchor=MSO_ANCHOR.MIDDLE)
    if page_no is not None:
        add_text(slide, Inches(8.0), Inches(7.14), Inches(1.2), Inches(0.32),
                 str(page_no), size=9, color=STEEL, anchor=MSO_ANCHOR.MIDDLE,
                 align=PP_ALIGN.CENTER)
    # footer logo (small) bottom-right on every slide
    place_logo(slide, Inches(11.7), Inches(7.02), Inches(1.5), Inches(0.42),
               compact=True)


def header_bar(slide, title, subtitle=None):
    rrect(slide, 0, 0, SW, Inches(1.02), fill=WF_DARK_RED, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    rrect(slide, 0, Inches(1.02), SW, Inches(0.07), fill=WF_GOLD, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, Inches(0.5), Inches(0.12), Inches(11.0), Inches(0.8),
             title, size=26, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.52), Inches(0.72), Inches(11.0), Inches(0.3),
                 subtitle, size=12, color=WF_GOLD, anchor=MSO_ANCHOR.MIDDLE)


def img_size(path):
    return Image.open(path).size


def add_image_fit(slide, path, x, y, w, h, align="center"):
    iw, ih = img_size(path)
    ar = iw / ih
    if w / h > ar:
        nh = h; nw = int(h * ar)
    else:
        nw = w; nh = int(w / ar)
    if align == "center":
        px = int(x + (w - nw) / 2)
    elif align == "left":
        px = int(x)
    else:
        px = int(x + (w - nw))
    py = int(y + (h - nh) / 2)
    slide.shapes.add_picture(path, px, py, width=nw, height=nh)
    return px, py, nw, nh


# ---------------------------------------------------------------------------
# Reusable content visuals
# ---------------------------------------------------------------------------
def before_after(slide, x, y, w, h, before_title, before_lines,
                 after_title, after_lines):
    bw = (w - Inches(1.0)) / 2
    # before (legacy)
    sp = rrect(slide, x, y, bw, h, fill=RGBColor(0xED, 0xF0, 0xF3),
               line=STEEL, line_w=1.5)
    add_text(slide, x, y + Inches(0.08), bw, Inches(0.35), before_title,
             size=13, color=STEEL, bold=True, align=PP_ALIGN.CENTER)
    add_bullets(slide, x + Inches(0.2), y + Inches(0.55), bw - Inches(0.4),
                h - Inches(0.65), before_lines, size=12, bullet_color=STEEL,
                color=INK)
    # arrow
    ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                x + bw + Inches(0.18), y + h/2 - Inches(0.22),
                                Inches(0.64), Inches(0.44))
    ar.fill.solid(); ar.fill.fore_color.rgb = WF_RED
    ar.line.fill.background(); ar.shadow.inherit = False
    # after (modern)
    ax0 = x + bw + Inches(1.0)
    rrect(slide, ax0, y, bw, h, fill=RGBColor(0xEA, 0xF2, 0xFB), line=BLUE,
          line_w=1.5)
    add_text(slide, ax0, y + Inches(0.08), bw, Inches(0.35), after_title,
             size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_bullets(slide, ax0 + Inches(0.2), y + Inches(0.55), bw - Inches(0.4),
                h - Inches(0.65), after_lines, size=12, bullet_color=BLUE,
                color=INK)


def wf_screen_mock(slide, x, y, w, h, screen_title, fields, badge_text,
                   badge_color):
    """Draw a Wells Fargo-themed React screen mockup with python-pptx shapes."""
    rrect(slide, x, y, w, h, fill=WHITE, line=BORDER, line_w=1.5)
    # dark-red header bar with logo
    rrect(slide, x, y, w, Inches(0.5), fill=WF_DARK_RED, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    place_logo(slide, x + Inches(0.1), y + Inches(0.09), Inches(1.15),
               Inches(0.32), compact=True)
    add_text(slide, x + Inches(1.5), y, w - Inches(1.6), Inches(0.5),
             screen_title, size=12, color=WHITE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    # field rows
    fy = y + Inches(0.7)
    for label, val in fields:
        add_text(slide, x + Inches(0.25), fy, Inches(1.7), Inches(0.3), label,
                 size=11, color=SLATE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        box = rrect(slide, x + Inches(2.0), fy, w - Inches(2.3), Inches(0.32),
                    fill=LIGHT, line=BORDER, line_w=0.75)
        shape_text(box, val, size=10.5, color=INK, bold=False,
                   align=PP_ALIGN.LEFT)
        fy += Inches(0.45)
    # status badge
    badge = rrect(slide, x + Inches(2.0), fy + Inches(0.02), Inches(1.4),
                  Inches(0.34), fill=badge_color, line=None)
    shape_text(badge, badge_text, size=11, color=WHITE, bold=True)
    add_text(slide, x + Inches(0.25), fy, Inches(1.7), Inches(0.34), "STATUS",
             size=11, color=SLATE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # gold accent action button
    btn = rrect(slide, x + w - Inches(1.7), y + h - Inches(0.55), Inches(1.4),
                Inches(0.38), fill=WF_GOLD, line=None)
    shape_text(btn, "Submit", size=11, color=INK, bold=True)


def dashboard_mock(slide, x, y, w, h):
    """Wells Fargo-themed reconciliation dashboard mockup."""
    rrect(slide, x, y, w, h, fill=WHITE, line=BORDER, line_w=1.5)
    rrect(slide, x, y, w, Inches(0.5), fill=WF_DARK_RED, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    place_logo(slide, x + Inches(0.1), y + Inches(0.09), Inches(1.15),
               Inches(0.32), compact=True)
    add_text(slide, x + Inches(1.5), y, w - Inches(1.6), Inches(0.5),
             "Settlement Reconciliation Dashboard", size=12, color=WHITE,
             bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # KPI cards
    kpis = [("Matched", "18,204", GREEN), ("Disputed", "126", WF_GOLD),
            ("Unmatched", "342", WF_RED)]
    cw = (w - Inches(0.8)) / 3
    kx = x + Inches(0.25)
    for name, val, col in kpis:
        card = rrect(slide, kx, y + Inches(0.65), cw - Inches(0.1),
                     Inches(0.9), fill=LIGHT, line=BORDER, line_w=1.0)
        add_text(slide, kx, y + Inches(0.72), cw - Inches(0.1), Inches(0.3),
                 name, size=11, color=SLATE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, kx, y + Inches(1.0), cw - Inches(0.1), Inches(0.45),
                 val, size=20, color=col, bold=True, align=PP_ALIGN.CENTER)
        kx += cw
    # per-network table
    ty = y + Inches(1.75)
    rows = [("Network", "Matched", "Disputed", "Unmatched"),
            ("VISA", "9,880", "54", "121"),
            ("MC", "6,102", "48", "150"),
            ("STAR", "2,222", "24", "71")]
    rowh = (y + h - Inches(0.2) - ty) / len(rows)
    colw = (w - Inches(0.5)) / 4
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            fill = WF_DARK_RED if ri == 0 else (LIGHT if ri % 2 else WHITE)
            tc = WHITE if ri == 0 else INK
            cx = x + Inches(0.25) + ci * colw
            cellsp = rrect(slide, cx, ty + ri * rowh, colw, rowh, fill=fill,
                           line=BORDER, line_w=0.5)
            shape_text(cellsp, cell, size=10, color=tc,
                       bold=(ri == 0 or ci == 0))


def draw_vicon(slide, cx, cy, size, kind, fg=WF_RED):
    """Draw a simple, renderer-safe vector icon centred at (cx, cy)."""
    s = size
    def oval(dx, dy, w, h, fill, line=None, lw=1.0):
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx + dx), int(cy + dy),
                                    int(w), int(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(lw)
        sp.shadow.inherit = False
        return sp
    def box(dx, dy, w, h, fill, line=None, lw=1.0, shp=MSO_SHAPE.ROUNDED_RECTANGLE):
        sp = slide.shapes.add_shape(shp, int(cx + dx), int(cy + dy), int(w),
                                    int(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(lw)
        sp.shadow.inherit = False
        return sp

    if kind == "money":
        c = oval(-s/2, -s/2, s, s, fg)
        shape_text(c, "$", size=int(s/12700*0.55), color=WHITE, bold=True)
    elif kind == "offsets":
        n = 3; bh = s*0.2; gap = s*0.14
        top = -(n*bh + (n-1)*gap)/2
        for i in range(n):
            box(-s/2, top + i*(bh+gap), s, bh, fg)
    elif kind == "stateless":
        box(-s/2, -s*0.22, s, s*0.44, fg, shp=MSO_SHAPE.LEFT_RIGHT_ARROW)
    elif kind == "gating":
        box(-s/2, -s/2, s, s, fg, shp=MSO_SHAPE.NO_SYMBOL)
    elif kind == "reconcile":
        c = oval(-s/2, -s/2, s, s, fg)
        shape_text(c, "=", size=int(s/12700*0.55), color=WHITE, bold=True)
    elif kind == "loop":
        box(-s/2, -s/2, s, s, fg, shp=MSO_SHAPE.CIRCULAR_ARROW)
    elif kind == "dto":
        box(-s*0.5, -s*0.15, s*0.62, s*0.62, None, line=fg, lw=2.2)
        box(-s*0.1, -s*0.45, s*0.62, s*0.62, fg)
    elif kind == "graph":
        r = s*0.24
        xs = [-s*0.5, 0, s*0.5]
        for i in range(len(xs) - 1):
            box(xs[i], -s*0.05, xs[i+1]-xs[i], s*0.1, fg,
                shp=MSO_SHAPE.RECTANGLE)
        for x in xs:
            oval(x - r/2, -r/2, r, r, fg)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # 1. TITLE ---------------------------------------------------------------
    s = _blank(prs)
    rrect(s, 0, 0, SW, SH, fill=WHITE, line=None, shape=MSO_SHAPE.RECTANGLE)
    rrect(s, 0, 0, SW, Inches(0.32), fill=WF_RED, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    rrect(s, 0, Inches(0.32), SW, Inches(0.10), fill=WF_GOLD, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    place_logo(s, Inches(4.67), Inches(0.9), Inches(4.0), Inches(1.2),
               placeholder_label="TITLE LOGO")
    add_text(s, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.4),
             "Accelerating Mainframe Modernization with Devin", size=40,
             color=INK, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.5), Inches(3.95), Inches(10.3), Inches(0.6),
             "Plastic Issuance & Settlement \u2014 an autonomous build-and-migrate demo",
             size=18, color=WF_DARK_RED, align=PP_ALIGN.CENTER)
    rrect(s, Inches(5.17), Inches(4.75), Inches(3.0), Inches(0.05),
          fill=WF_GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4),
             "Wells Fargo  +  Mphasis  \u00b7  The Next Applied", size=13,
             color=STEEL, align=PP_ALIGN.CENTER)
    place_logo(s, Inches(11.55), Inches(6.98), Inches(1.65), Inches(0.5))

    # 2. AGENDA --------------------------------------------------------------
    s = _blank(prs); header_bar(s, "Agenda")
    agenda = [
        "What Devin is and how its delivery loop works",
        "Why Devin is more efficient than autocomplete assistants",
        "Valuable Devin capabilities",
        "How we apply them to this migration",
        "The demo application architecture",
        "Act 1 \u2014 build the mainframe app",
        "Act 2 \u2014 migrate to .NET + React + PostgreSQL",
        "Business-language prompting, precautions & safe shortcuts",
        "Demo runbook",
    ]
    y = Inches(1.35)
    for i, item in enumerate(agenda):
        num = rrect(s, Inches(0.7), y, Inches(0.55), Inches(0.5),
                    fill=WF_RED if i % 2 == 0 else WF_DARK_RED, line=None)
        shape_text(num, str(i + 1), size=16, color=WHITE, bold=True)
        card = rrect(s, Inches(1.4), y, Inches(10.9), Inches(0.5),
                     fill=LIGHT, line=BORDER, line_w=1.0)
        shape_text(card, item, size=15, color=INK, bold=False,
                   align=PP_ALIGN.LEFT)
        y += Inches(0.62)
    add_footer(s, 2)

    # 3. WHAT IS DEVIN -------------------------------------------------------
    s = _blank(prs); header_bar(s, "What Is Devin?")
    add_text(s, Inches(0.6), Inches(1.25), Inches(12.1), Inches(0.7),
             "Devin is an autonomous AI software engineer: give it a goal and it "
             "plans, writes and runs code on its own machine, verifies the results, "
             "and opens a pull request \u2014 iterating until the work is correct.",
             size=17, color=SLATE)
    add_image_fit(s, os.path.join(IMG, "devin-workflow-loop.png"),
                  Inches(1.6), Inches(2.15), Inches(10.1), Inches(4.7))
    add_footer(s, 3)

    # 4. WHY MORE EFFICIENT --------------------------------------------------
    s = _blank(prs); header_bar(s, "Why Devin Is More Efficient")
    add_image_fit(s, os.path.join(IMG, "devin-vs-others.png"),
                  Inches(0.5), Inches(1.25), Inches(8.4), Inches(5.7))
    add_bullets(s, Inches(9.05), Inches(1.6), Inches(4.0), Inches(5.0), [
        "Owns whole tasks, not just the next line",
        "Runs, tests and verifies its own work",
        "Closes the loop by diffing real outputs",
        "Runs many scoped sessions in parallel",
    ], size=15)
    add_footer(s, 4)

    # 5. VALUABLE FEATURES ---------------------------------------------------
    s = _blank(prs); header_bar(s, "Valuable Devin Capabilities")
    add_image_fit(s, os.path.join(IMG, "devin-capability-grid.png"),
                  Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.5))
    add_footer(s, 5)

    # 6. HOW WE USE THEM -----------------------------------------------------
    s = _blank(prs); header_bar(s, "How We Use These Features Here")
    add_image_fit(s, os.path.join(IMG, "devin-parallel-agents.png"),
                  Inches(0.4), Inches(1.2), Inches(7.4), Inches(5.7))
    add_bullets(s, Inches(7.95), Inches(1.45), Inches(5.1), Inches(5.4), [
        "Scoped sessions per program area (PI vs. Settlement)",
        "Verification via LOCAL/run.sh output diffs",
        "Knowledge onboarding from README.md & DOC/RUN_GUIDE.md",
        "Parallel PI and Settlement sessions",
        "Copybook \u2192 DTO automation",
        "CI acts as the regression oracle",
    ], size=14)
    add_footer(s, 6)

    # 7. ARCHITECTURE --------------------------------------------------------
    s = _blank(prs); header_bar(s, "Demo Application Architecture")
    add_image_fit(s, os.path.join(IMG, "architecture-diagram.png"),
                  Inches(0.5), Inches(1.25), Inches(12.3), Inches(5.7))
    add_footer(s, 7)

    # 8. ACT 1 ---------------------------------------------------------------
    s = _blank(prs); header_bar(s, "Act 1 \u2014 Build the Mainframe App",
                                "Devin builds and runs the full pipeline locally")
    add_image_fit(s, os.path.join(IMG, "pipeline-flow.png"),
                  Inches(0.5), Inches(1.35), Inches(12.3), Inches(5.5))
    add_footer(s, 8)

    # 9. ACT 2: DATA STORE ---------------------------------------------------
    s = _blank(prs)
    header_bar(s, "Act 2 \u2014 Data Store \u2192 PostgreSQL + .NET",
               "Exact decimal precision \u2014 no float/double drift")
    before_after(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(3.9),
                 "MAINFRAME",
                 ["DB2 relational tables", "VSAM KSDS files",
                  "Packed-decimal money fields", "Fixed record layouts"],
                 "MODERN",
                 [".NET data services", "PostgreSQL",
                  "DECIMAL columns (exact to the cent)",
                  "Row-count & total verification vs. source"])
    add_text(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(1.2),
             "Precaution: amounts stored as DECIMAL, never float/double. Devin "
             "reconciles row counts and totals to the cent after migration.",
             size=14, color=WF_DARK_RED, bold=True)
    add_footer(s, 9)

    # 10. ACT 2: BATCH -------------------------------------------------------
    s = _blank(prs)
    header_bar(s, "Act 2 \u2014 Batch Job \u2192 .NET Background Service",
               "Prove output equivalence line-for-line")
    before_after(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(3.9),
                 "MAINFRAME",
                 ["COBOL batch programs", "JCL job steps & COND codes",
                  "SYSOUT reports", "Sequential read-process-write"],
                 "MODERN",
                 [".NET background services", "Same validation & rejection rules",
                  "Same summary reports", "Outputs diffed against baseline"])
    add_text(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(1.2),
             "The existing OUTPUT/*.txt reports are the reference oracle: the new "
             "service must reproduce them exactly on the same input.",
             size=14, color=WF_DARK_RED, bold=True)
    add_footer(s, 10)

    # 11. ACT 2: CICS -> REST + REACT ---------------------------------------
    s = _blank(prs)
    header_bar(s, "Act 2 \u2014 CICS Green Screen \u2192 REST + React",
               "Stateless APIs \u00b7 Wells Fargo-themed UI")
    before_after(s, Inches(0.6), Inches(1.35), Inches(6.0), Inches(3.0),
                 "MAINFRAME",
                 ["CICS 3270 green screens", "Pseudo-conversational COMMAREA state"],
                 "MODERN",
                 ["Stateless REST endpoints", "React screens on a shared theme"])
    wf_screen_mock(s, Inches(7.0), Inches(1.35), Inches(5.7), Inches(3.7),
                   "Card Inquiry",
                   [("Card Number", "4000 12** **** 0017"),
                    ("Customer", "J. RIVERA"),
                    ("Card Type", "Debit")],
                   "ACTIVE", GREEN)
    add_bullets(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(2.1), [
        "React screens are built from a shared Wells Fargo design system / component "
        "library (ThemeProvider, AppHeader with logo, Button/Card/Table/StatusBadge)",
        "Status shown as coloured badges: green = Active, red = Blocked, grey = Closed",
    ], size=13)
    add_footer(s, 11)

    # 12. ACT 2: SETTLEMENT DASHBOARD ---------------------------------------
    s = _blank(prs)
    header_bar(s, "Act 2 \u2014 Settlement \u2192 Service + Dashboard",
               "Wells Fargo-themed reconciliation dashboard")
    before_after(s, Inches(0.6), Inches(1.35), Inches(5.4), Inches(4.3),
                 "MAINFRAME",
                 ["Batch matching over VSAM", "Easytrieve / SAS reports",
                  "Per-network reconciliation totals"],
                 "MODERN",
                 [".NET settlement service", "React dashboard",
                  "Same matched/disputed/unmatched totals"])
    dashboard_mock(s, Inches(6.3), Inches(1.35), Inches(6.4), Inches(4.3))
    add_text(s, Inches(0.6), Inches(5.85), Inches(12.1), Inches(1.0),
             "Matched / disputed / unmatched totals per network must equal today's "
             "reconciliation output before the legacy job is retired.",
             size=13, color=WF_DARK_RED, bold=True)
    add_footer(s, 12)

    # 13. BUSINESS PROMPT TECHNIQUES ----------------------------------------
    s = _blank(prs); header_bar(s, "Business-Language Prompting",
                                "Plain functional prompts \u2014 no code or file names")
    prompts = [
        ("A \u00b7 Issue cards with an industry-standard card-number checksum; "
         "reject-and-report bad records, don't fail the run.",
         "Points at Luhn without naming it; captures batch resilience."),
        ("A \u00b7 Match daily settlement into matched / disputed / unmatched, "
         "efficiently for large files.",
         "Yields the sorted-merge algorithm and per-network totals."),
        ("B \u00b7 Migrate one business area at a time, run in parallel, prove "
         "equivalence; rules unchanged.",
         "Sets a safe, verifiable incremental strategy."),
        ("B \u00b7 Store money exactly \u2014 no floating point; verify totals to "
         "the cent.",
         "Locks in decimal precision as an acceptance test."),
        ("B \u00b7 Build a reusable Wells Fargo React design system first "
         "(theme, header, badges).",
         "Every screen inherits one brand-accurate look."),
    ]
    rows = len(prompts) + 1
    tbl = s.shapes.add_table(rows, 2, Inches(0.6), Inches(1.5), Inches(12.1),
                             Inches(5.2)).table
    tbl.columns[0].width = Inches(7.6)
    tbl.columns[1].width = Inches(4.5)
    hdr = ["Business-language prompt", "Why it works"]
    for c, txt in enumerate(hdr):
        cell = tbl.cell(0, c); cell.text = txt
        cell.fill.solid(); cell.fill.fore_color.rgb = WF_DARK_RED
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(13); p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = WHITE
    for r, (pr, why) in enumerate(prompts, start=1):
        for c, txt in enumerate((pr, why)):
            cell = tbl.cell(r, c); cell.text = txt
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.color.rgb = INK if c == 0 else WF_DARK_RED
            p.runs[0].font.bold = (c == 1)
    add_footer(s, 13)

    # 14. PRECAUTIONS --------------------------------------------------------
    s = _blank(prs); header_bar(s, "Precautions")
    precautions = [
        ("money", "Money & decimal precision", "Use exact DECIMAL; never float/double."),
        ("offsets", "Fixed-width file offsets", "Preserve byte positions when parsing records."),
        ("stateless", "Stateless APIs", "No hidden session state between screens."),
        ("gating", "Job-failure dependency gating", "Never run a step if the prior one failed."),
        ("reconcile", "Parallel-run reconciliation", "Compare old vs. new outputs before cutover."),
    ]
    cw = Inches(2.35); gap = Inches(0.14)
    total = 5 * cw + 4 * gap
    x0 = (SW - total) / 2
    for i, (kind, title, desc) in enumerate(precautions):
        x = x0 + i * (cw + gap)
        card = rrect(s, x, Inches(1.9), cw, Inches(3.4), fill=WHITE,
                     line=BORDER, line_w=1.5)
        draw_vicon(s, x + cw/2, Inches(2.6), Inches(0.8), kind, fg=WF_RED)
        rrect(s, x + Inches(0.3), Inches(3.15), cw - Inches(0.6), Inches(0.04),
              fill=WF_GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, x + Inches(0.1), Inches(3.3), cw - Inches(0.2), Inches(0.8),
                 title, size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(4.1), cw - Inches(0.2), Inches(1.1),
                 desc, size=11, color=SLATE, align=PP_ALIGN.CENTER)
    add_footer(s, 14)

    # 15. SAFE SHORTCUTS -----------------------------------------------------
    s = _blank(prs); header_bar(s, "Safe Shortcuts & Automations")
    shortcuts = [
        ("loop", "Local pipeline as CI regression oracle",
         "LOCAL/build.sh + run.sh give a laptop-runnable baseline; wire it into CI "
         "so every change is diffed against known-good outputs."),
        ("dto", "Layout-to-DTO generation",
         "Turn fixed-width record layouts into typed DTOs automatically \u2014 "
         "preserving field offsets and decimal scale."),
        ("graph", "Schedule dependency-graph extraction",
         "Read the batch schedule and derive the success-gated dependency graph "
         "the orchestration must preserve."),
    ]
    y = Inches(1.6)
    for kind, title, desc in shortcuts:
        card = rrect(s, Inches(0.7), y, Inches(11.9), Inches(1.5), fill=LIGHT,
                     line=BORDER, line_w=1.0)
        ic = rrect(s, Inches(0.9), y + Inches(0.3), Inches(0.9), Inches(0.9),
                   fill=WF_RED, line=None)
        draw_vicon(s, Inches(0.9) + Inches(0.45), y + Inches(0.3) + Inches(0.45),
                   Inches(0.5), kind, fg=WHITE)
        add_text(s, Inches(2.05), y + Inches(0.18), Inches(10.3), Inches(0.5),
                 title, size=16, color=INK, bold=True)
        add_text(s, Inches(2.05), y + Inches(0.68), Inches(10.3), Inches(0.75),
                 desc, size=12.5, color=SLATE)
        y += Inches(1.7)
    add_footer(s, 15)

    # 16. RUNBOOK ------------------------------------------------------------
    s = _blank(prs); header_bar(s, "Demo Runbook",
                                "Exact commands for the local end-to-end demo")
    steps = [
        ("1  Build all programs", "cd LOCAL && bash build.sh"),
        ("2  Run the full pipeline", "cd LOCAL && bash run.sh"),
        ("3  Inspect key outputs", "ls OUTPUT/   # CARDRPT.txt, STLRPT*.txt, MGTRPT.txt, SASOUT.csv"),
        ("4  Prove equivalence (old vs. new)", "diff OUTPUT/CARDRPT.txt  NEW/OUTPUT/CARDRPT.txt"),
    ]
    y = Inches(1.5)
    for label, cmd in steps:
        add_text(s, Inches(0.7), y, Inches(11.9), Inches(0.35), label, size=15,
                 color=WF_DARK_RED, bold=True)
        code = rrect(s, Inches(0.7), y + Inches(0.38), Inches(11.9),
                     Inches(0.6), fill=INK, line=None)
        tf = code.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(10)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = "$ " + cmd
        r.font.size = Pt(14); r.font.name = "Consolas"
        r.font.color.rgb = WF_GOLD; r.font.bold = True
        y += Inches(1.18)
    add_footer(s, 16)

    # 17. CLOSING ------------------------------------------------------------
    s = _blank(prs)
    rrect(s, 0, 0, SW, SH, fill=WF_DARK_RED, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    rrect(s, 0, Inches(3.0), SW, Inches(0.08), fill=WF_GOLD, line=None,
          shape=MSO_SHAPE.RECTANGLE)
    add_text(s, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.2),
             "Thank You", size=48, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.8), Inches(3.2), Inches(11.7), Inches(0.8),
             "Questions & Discussion", size=24, color=WF_GOLD,
             align=PP_ALIGN.CENTER)
    place_logo(s, Inches(4.67), Inches(4.4), Inches(4.0), Inches(1.2),
               placeholder_label="LOGO")
    add_text(s, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.4),
             "Wells Fargo  +  Mphasis  \u00b7  The Next Applied", size=13,
             color=WHITE, align=PP_ALIGN.CENTER)

    prs.save(OUT_PPTX)
    print("Saved", OUT_PPTX)
    if LOGO:
        print("Logo used:", LOGO)
    else:
        print("\n" + "=" * 68)
        print("NOTE: No logo image was found. A clearly-marked placeholder was")
        print("inserted on the title slide, every footer, and the closing slide.")
        print("To brand the deck, drop the pasted Wells Fargo + Mphasis logo at:")
        for name in LOGO_CANDIDATES[:3]:
            print("   " + os.path.join(LOGO_DIR, name))
        print("then re-run:  python3 tools/demo/gen_pptx.py")
        print("=" * 68)


if __name__ == "__main__":
    build()
